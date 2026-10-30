"""Planejamento anual por meses em Word (.docx), versão editável da área premium.

Usa o mesmo calendário e a mesma distribuição de _fonte/planejamento/gerar.py,
então o Word e o PDF gratuito sempre batem.

Uso (na raiz do repositório):
    python3 _fonte/premium/editaveis/planejamento_docx.py _privado/premium/arquivos/planejamento
Autora: Júlia Roberta Quoos Alves.
"""
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
sys.path.insert(0, os.path.join(RAIZ, '_fonte', 'planejamento'))
import gerar as G  # noqa: E402

AUTORA = 'Júlia Roberta Quoos Alves'
VERDE, AMARELO, CINZA = 'E3F1E8', 'FFF9C4', 'F4F4F2'


def fundo(celula, cor):
    pr = celula._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), cor)
    pr.append(sh)


def repetir_cabecalho(linha):
    pr = linha._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); pr.append(el)


def link(par, texto, url):
    rid = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '16'); rpr.append(sz)
    r.append(rpr); t = OxmlElement('w:t'); t.text = texto; r.append(t); h.append(r)
    par._p.append(h)


def texto(celula, t, negrito=False, tam=8.5, cor=None):
    p = celula.paragraphs[0] if not celula.paragraphs[0].text else celula.add_paragraph()
    run = p.add_run(t); run.bold = negrito; run.font.size = Pt(tam)
    if cor:
        run.font.color.rgb = RGBColor.from_string(cor)
    return p


def documento():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for lado in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, lado, Mm(12))
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(9.5)
    for nome, tam in (('Heading 1', 20), ('Heading 2', 12)):  # sem o azul do modelo padrão do Word
        doc.styles[nome].font.color.rgb = RGBColor.from_string('111111'); doc.styles[nome].font.size = Pt(tam)
    doc.core_properties.author = AUTORA
    rod = sec.footer.paragraphs[0]
    rod.text = f'{AUTORA} · ohtopico.com.br · versão editável para apoiadores do Catarse'
    rod.runs[0].font.size = Pt(8)
    return doc


def gerar(comp, dados, destino):
    c = G.COMP[comp]
    sem = G.semanas()
    total = len(sem)
    meses = []
    for s in sem:
        if not meses or meses[-1][0] != s['seg'].month:
            meses.append((s['seg'].month, []))
        meses[-1][1].append(s)
    doc = documento()
    doc.core_properties.title = f'Planejamento anual de {c["nome"]} por meses · {G.ANO}'
    doc.add_heading(f'Planejamento anual de {c["nome"]} por meses · {G.ANO}', 1).runs[0].font.size = Pt(26)
    doc.add_paragraph(f'6º ao 9º ano do Ensino Fundamental. {len(dados)} habilidades da BNCC distribuídas semana a semana, '
                      f'com 1 aula por semana por turma, de {G.d(G.INICIO)} a {G.d(G.FIM)}/{G.ANO} ({total} semanas letivas).')
    doc.add_heading('Como adaptar este arquivo', 2)
    for t in ['Troque as datas da coluna Semana pelas do calendário da sua rede ou escola.',
              'A coluna Observações está em branco para as suas anotações, turmas e ajustes.',
              'Se a aula cair em feriado, passe o conteúdo para a semana seguinte. A semana de revisão do bimestre absorve o atraso.',
              f'{G.RECESSO[2]}: {G.d(G.RECESSO[0])} a {G.d(G.RECESSO[1])}. Bimestres a partir de '
              + ', '.join(G.d(b) for b in G.BIMESTRES) + '.',
              'Feriados considerados: ' + '; '.join(f'{G.d(f)} ({G.DIAS[f.weekday()]}) {n}' for f, n in G.FERIADOS) + '.']:
        doc.add_paragraph(t, style='List Bullet')
    doc.add_paragraph(f'Fonte: {c["fonte"]}').runs[0].font.size = Pt(8.5)
    doc.add_paragraph(f'Material criado por {AUTORA}. Uso livre em sala de aula, com citação da fonte. '
                      'Arquivo exclusivo para apoiadores do Catarse: por favor, não redistribua.').runs[0].font.size = Pt(8.5)

    larg = [Mm(30), Mm(88), Mm(48), Mm(64), Mm(43)]
    for ano in (6, 7, 8, 9):
        habs = [h for h in dados if h['ano'] == ano]
        plano = G.plano_ano(habs, sem)
        doc.add_page_break()
        doc.add_heading(f'{ano}º ano · {c["nome"]}', 1)
        doc.add_paragraph(f'{len(habs)} habilidades em {total} aulas.')
        for m, ss in meses:
            bims = sorted({s['bim'] for s in ss})
            doc.add_heading(f'{G.MESES[m]} · {len(ss)} aulas · ' + ' e '.join(f'{b}º' for b in bims) + ' bimestre', 2)
            tab = doc.add_table(rows=1, cols=5)
            tab.style = 'Table Grid'; tab.alignment = WD_TABLE_ALIGNMENT.CENTER; tab.autofit = False
            for i, nome in enumerate(['Semana', 'Habilidade da BNCC', 'Foco da aula', 'Materiais do ohtopico', 'Observações']):
                cel = tab.rows[0].cells[i]; cel.width = larg[i]
                texto(cel, nome, negrito=True, tam=8.5); fundo(cel, VERDE)
            repetir_cabecalho(tab.rows[0])

            def faixa(t):
                cel = tab.add_row().cells
                junta = cel[0].merge(cel[4]); texto(junta, t, negrito=True); fundo(junta, AMARELO)

            for s, itens, esp in plano:
                if s['seg'].month != m:
                    continue
                n = sem.index(s) + 1
                if m == 7 and s['seg'] > G.RECESSO[1] and sem[n - 2]['seg'] < G.RECESSO[0]:
                    faixa(f'{G.RECESSO[2]}: {G.d(G.RECESSO[0])} a {G.d(G.RECESSO[1])}')
                if s['bim'] > 1 and s['seg'] == G.BIMESTRES[s['bim'] - 1]:
                    faixa(f'Início do {s["bim"]}º bimestre')
                cel = tab.add_row().cells
                for i, w in enumerate(larg):
                    cel[i].width = w
                texto(cel[0], f'Aula {n}', negrito=True)
                texto(cel[0], f'{G.d(s["seg"])} a {G.d(s["sex"])}')
                for f, nm in s['fer']:
                    texto(cel[0], f'Feriado {G.DIAS[f.weekday()]} {G.d(f)}: {nm}', tam=7.5, cor='8A1328')
                if esp:
                    junta = cel[1].merge(cel[2])
                    texto(junta, esp[0], negrito=True); texto(junta, esp[1])
                    for x in cel:
                        fundo(x, CINZA)
                    continue
                for h, _, _, k, _ in itens:
                    p = texto(cel[1], h['codigo'] + ' ', negrito=True)
                    p.add_run(h['texto'] if k == 1 else f'continuação · {h["objeto"]}').font.size = Pt(8.5)
                h, foco, ms, k, q = itens[0]
                if len(itens) > 1:
                    texto(cel[2], f'{len(itens)} habilidades na mesma aula', negrito=True)
                    texto(cel[2], 'Trabalhe as duas a partir de um problema comum.', tam=8)
                    ms = [m2 for it in itens for m2 in it[2]]
                else:
                    texto(cel[2], foco, negrito=True)
                    texto(cel[2], f'Aula {k} de {q} desta habilidade' if q > 1 else 'Uma aula para esta habilidade', tam=8)
                ms = list({x['u']: x for x in ms}.values())
                if not ms:
                    texto(cel[3], 'Livro didático ou material próprio', tam=8)
                for x in ms:
                    p = cel[3].paragraphs[0] if not cel[3].paragraphs[0].text and not cel[3].paragraphs[0]._p.findall(qn('w:hyperlink')) else cel[3].add_paragraph()
                    link(p, x['t'], x['u'])
    nome = os.path.join(destino, f'planejamento-mensal-{c["arquivo"]}.docx')
    doc.save(nome)
    return nome


def main():
    destino = sys.argv[1] if len(sys.argv) > 1 else '_privado/premium/arquivos/planejamento'
    os.makedirs(destino, exist_ok=True)
    dados = json.load(open(os.path.join(RAIZ, '_fonte', 'planejamento', 'habilidades.json'), encoding='utf-8'))
    for comp in G.COMP:
        print('ok', gerar(comp, dados[comp], destino))


if __name__ == '__main__':
    main()
