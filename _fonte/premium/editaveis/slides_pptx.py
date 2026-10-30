"""Slides das 42 aulas da BNCC Computação em PowerPoint (.pptx) editável, para a área premium.

Lê as aulas do gerador dos slides em PDF (aulas6.py a aulas9.py, lib.py e svg.py) e monta
cada slide com caixas de texto editáveis. Os diagramas entram como imagem PNG.
Abre no PowerPoint, no Google Apresentações e no Canva (Importar arquivo).

Uso:
    python3 _fonte/premium/editaveis/slides_pptx.py --gerador CAMINHO/bncc-computacao/gerador \
        --fontes CAMINHO/fontes --saida _privado/premium/arquivos/slides
--fontes é a pasta com Lexend-400.ttf, Lexend-600.ttf, Sniglet-800.ttf e JetBrainsMono-Regular.ttf,
usadas só para desenhar os diagramas (precisa do Node com Playwright).
Autora: Júlia Roberta Quoos Alves.
"""
import argparse
import importlib
import json
import math
import os
import re
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

AUTORA = 'Júlia Roberta Quoos Alves'
TINTA, TINTA2, PAPEL, PAPEL2 = '111111', '444444', 'FFFFFF', 'F4F4F2'
MARCA, OK, MAL, VERDE = 'FFF36B', 'E3F4E8', 'FBE6EA', '157F3B'
CORES = {'': OK, 'y': MARCA, 'r': MAL, 'w': PAPEL}
TIT, CORPO, MONO = 'Sniglet', 'Lexend', 'JetBrains Mono'
PX = 12700 * 2 / 3  # o molde em HTML tem 1440 x 810 px; 1 px = 2/3 pt
CX, CY, CW, CH = 128, 150, 1172, 530  # área útil dentro da janela
LARG_TIT = 0.84  # largura média de uma letra maiúscula da Sniglet, em relação ao tamanho da fonte


def e(px):
    return Emu(int(px * PX))


def cor(h):
    return RGBColor.from_string(h)


def runs(s):
    """'texto **negrito** [[destaque]]' -> [(texto, negrito, destaque)]."""
    out = []
    for parte in re.split(r'(\*\*.+?\*\*|\[\[.+?\]\])', s.replace('<br>', '\n')):
        if not parte:
            continue
        if parte.startswith('**'):
            out.append((parte[2:-2], True, False))
        elif parte.startswith('[['):
            out.append((parte[2:-2], False, True))
        else:
            out.append((parte, False, False))
    return out


def limpo(s):
    return re.sub(r'\[\[|\]\]|\*\*', '', s.replace('<br>', ' '))


def caber(textos, w, h, fs, fs_min=14, alt=1.42, larg=0.56, extra=0):
    """Diminui a fonte (px) até o texto estimado caber na caixa."""
    while fs > fs_min:
        linhas = sum(max(1, math.ceil(len(limpo(t)) * fs * larg / w)) + t.count('<br>') for t in textos)
        if linhas * fs * alt + extra * fs <= h:
            break
        fs -= 1
    return fs


def caixa(sl, x, y, w, h, fundo=None, borda=True, raio=0.08, forma=MSO_SHAPE.ROUNDED_RECTANGLE, espessura=3):
    s = sl.shapes.add_shape(forma, e(x), e(y), e(w), e(h))
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = raio
    if fundo:
        s.fill.solid(); s.fill.fore_color.rgb = cor(fundo)
    else:
        s.fill.background()
    if borda:
        s.line.color.rgb = cor(TINTA); s.line.width = Pt(espessura * 2 / 3)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def escrever(alvo, pars, fs, fonte=CORPO, cor_txt=TINTA, alinhar=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP,
             caixa_alta=False, marcador=None, espaco=10, margem=0):
    """pars = lista de strings com ** e [[ ]]; marcador = None | 'dot' | 'num'."""
    tf = alvo.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    for lado in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, lado, e(margem))
    for i, par in enumerate(pars):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhar
        p.space_after = Pt(espaco * 2 / 3)
        p.line_spacing = 1.12
        if marcador:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(e(fs * 1.5)))); pPr.set('indent', str(int(-e(fs * 1.5))))
            if marcador == 'num':
                b = pPr.makeelement(qn('a:buAutoNum'), {'type': 'arabicPeriod'})
            else:
                cl = pPr.makeelement(qn('a:buClr'), {}); s = cl.makeelement(qn('a:srgbClr'), {'val': VERDE}); cl.append(s); pPr.append(cl)
                b = pPr.makeelement(qn('a:buChar'), {'char': '■'})
            pPr.append(b)
        for txt, neg, dest in runs(par):
            r = p.add_run()
            r.text = txt.upper() if caixa_alta else txt
            f = r.font
            f.size = Pt(fs * 2 / 3); f.name = fonte; f.bold = neg or fonte == TIT
            f.color.rgb = cor(cor_txt)
            if dest:
                rPr = r._r.get_or_add_rPr()
                hl = rPr.makeelement(qn('a:highlight'), {}); c = hl.makeelement(qn('a:srgbClr'), {'val': MARCA}); hl.append(c)
                rPr.append(hl)
    return tf


def texto(sl, x, y, w, h, pars, fs, **kw):
    tb = sl.shapes.add_textbox(e(x), e(y), e(w), e(h))
    escrever(tb, pars, fs, **kw)
    return tb


def titulo(sl, t, fs=50, y=CY, w=CW):
    fs = caber([t], w, fs * 2.6, fs, 30, alt=1.15, larg=LARG_TIT)
    linhas = max(1, math.ceil(len(limpo(t)) * fs * LARG_TIT / w))
    texto(sl, CX, y, w, linhas * fs * 1.25, [t], fs, fonte=TIT, caixa_alta=True, ancora=MSO_ANCHOR.TOP)
    return y + linhas * fs * 1.25 + 24


def molde(prs, aula, n, total, rotulo):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = cor(PAPEL2)
    caixa(sl, 70, 54, 1316, 680, TINTA, borda=False, raio=0.05)          # sombra
    caixa(sl, 56, 40, 1316, 680, PAPEL, raio=0.05, espessura=3.5)        # janela
    barra = caixa(sl, 58, 42, 1312, 56, PAPEL2, borda=False, forma=MSO_SHAPE.RECTANGLE)
    sl.shapes.add_connector(1, e(58), e(98), e(1370), e(98)).line.color.rgb = cor(TINTA)
    escrever(barra, [aula['cod'] + '.ppt'], 18, fonte=TIT, caixa_alta=True, ancora=MSO_ANCHOR.MIDDLE, margem=30)
    for i, s in enumerate(('–', '□', '✕')):
        b = caixa(sl, 1250 + i * 36, 56, 28, 28, PAPEL, forma=MSO_SHAPE.OVAL, espessura=2.5)
        escrever(b, [s], 13, alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    texto(sl, 60, 758, 700, 36, [AUTORA], 18, fonte=TIT, caixa_alta=True, ancora=MSO_ANCHOR.MIDDLE)
    texto(sl, 900, 758, 120, 36, [f'{n}/{total}'], 16, fonte=TIT, alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    pill = caixa(sl, 1180, 760, 190, 34, PAPEL, raio=0.5, espessura=2.5)
    escrever(pill, ['ohtopico.com.br'], 16, alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    sl.name = rotulo[:60]
    return sl


def lista(sl, y, itens, fs=31, marcador='dot', rod=None, x=CX, w=CW):
    h = CY + CH - y - (60 if rod else 0)
    fs = caber(itens, w - fs * 1.5, h, fs, 16, extra=0.5 * len(itens))
    texto(sl, x, y, w, h, itens, fs, marcador=marcador, espaco=14)
    if rod:
        texto(sl, x, CY + CH - 56, w, 56, [rod], min(fs, 24), cor_txt=TINTA2)


def paragrafos(sl, x, y, w, h, ps, fs=31, **kw):
    fs = caber(ps, w, h, fs, 16, extra=0.6 * len(ps))
    texto(sl, x, y, w, h, ps, fs, espaco=16, **kw)


def figura(sl, png, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(png).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    pic = sl.shapes.add_picture(png, e(x + (w - pw) / 2), e(y + (h - ph) / 2), e(pw), e(ph))
    return pic


# ---------------------------------------------------------------- tipos de slide

def s_capa(sl, a, d, ctx):
    from lib import NOMES_EIXO
    chip = caixa(sl, CX, CY, 520, 44, PAPEL, raio=0.5)
    escrever(chip, [f'BNCC Computação · {a["ano"]}º ano · [[{a["cod"]}]]'], 20, fonte=TIT, caixa_alta=True,
             alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    y = titulo(sl, a['titulo'], 74, y=CY + 80)
    texto(sl, CX, y, CW, 80, [a['sub']], 31, cor_txt=TINTA2)
    xs = CX
    for t, c in ((NOMES_EIXO[a['eixo']], OK), (limpo(a['objeto']), MARCA), ('Aula expositiva · 30 min', PAPEL)):
        w = len(t) * 12 + 50
        p = caixa(sl, xs, y + 120, w, 40, c, raio=0.5, espessura=2.5)
        escrever(p, [t], 20, alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
        xs += w + 14


def s_perg(sl, a, d, ctx):
    fs = caber([d['q']], CW, 330, 66, 36, alt=1.15, larg=LARG_TIT)
    texto(sl, CX, CY + 40, CW, 340, [d['q']], fs, fonte=TIT, caixa_alta=True, ancora=MSO_ANCHOR.BOTTOM)
    texto(sl, CX, CY + 400, CW, 90, [d.get('dica', 'Pense antes de passar o slide.')], 31, cor_txt=TINTA2)


def s_hab(sl, a, d, ctx):
    y = titulo(sl, 'O que vamos [[aprender]]')
    fs = caber([a['hab']], CW - 56, 150, 29, 18)
    linhas = max(1, math.ceil(len(a['hab']) * fs * 0.6 / (CW - 56)))
    hh = 90 + linhas * fs * 1.45
    b = caixa(sl, CX, y, CW, hh, OK, raio=0.12)
    tf = escrever(b, [f'**{a["cod"]}**   {limpo(a["objeto"]).upper()}', a['hab']], fs, margem=24, espaco=8)
    tf.paragraphs[0].runs[0].font.name = TIT
    y += hh + 22
    texto(sl, CX, y, CW, 40, ['**Ao final da aula, você vai conseguir:**'], 26)
    lista(sl, y + 46, a['objetivos'], 27)


def s_texto(sl, a, d, ctx):
    y = titulo(sl, d['t'])
    paragrafos(sl, CX, y, CW, CY + CH - y, d['ps'])


def s_lista(sl, a, d, ctx):
    y = titulo(sl, d['t'])
    lista(sl, y, d['itens'], marcador='num' if d.get('num') else 'dot', rod=d.get('rod'))


def s_cards(sl, a, d, ctx):
    y = titulo(sl, d['t'])
    cs = d['cards']
    cores = d.get('cores') or ['', 'y', 'r', 'w', '', 'y']
    cols = d.get('cols', len(cs) if len(cs) <= 3 else 2)
    linhas = math.ceil(len(cs) / cols)
    gap = 22
    rod = d.get('rod')
    hdisp = CY + CH - y - (64 if rod else 0)
    w = (CW - gap * (cols - 1)) / cols
    h = (hdisp - gap * (linhas - 1)) / linhas
    centro = d.get('c', True)
    for i, c in enumerate(cs):
        x = CX + (i % cols) * (w + gap); yy = y + (i // cols) * (h + gap)
        b = caixa(sl, x, yy, w, h, CORES[cores[i % len(cores)]], raio=0.12)
        pars, tams = [], []
        if c[0]:
            pars.append(c[0]); tams.append(40)
        if c[1]:
            pars.append(f'**{c[1]}**'); tams.append(28)
        pars.append(c[2]); tams.append(25)
        k = caber(pars, w - 40, h - 30, 25, 13)
        tf = escrever(b, pars, k, alinhar=PP_ALIGN.CENTER if centro else PP_ALIGN.LEFT, ancora=MSO_ANCHOR.MIDDLE, margem=18, espaco=4)
        for p, t in zip(tf.paragraphs, tams):
            for r in p.runs:
                r.font.size = Pt(max(13, t * k / 25) * 2 / 3)
                if t == 40:
                    r.font.name = TIT; r.font.bold = True
    if rod:
        texto(sl, CX, CY + CH - 58, CW, 58, [rod], 24, cor_txt=TINTA2, alinhar=PP_ALIGN.CENTER if centro else PP_ALIGN.LEFT)


def s_code(sl, a, d, ctx):
    y = titulo(sl, d['t'])
    lado = d.get('ps', [])
    w = CW * (1.25 / 2.25) - 22 if lado else CW
    linhas = d['code'].split('\n')
    fs = 27
    while fs > 14 and (len(linhas) * fs * 1.5 > CY + CH - y - 40 or max(len(l) for l in linhas) * fs * 0.6 > w - 56):
        fs -= 1
    hh = min(CY + CH - y, len(linhas) * fs * 1.5 + 48)
    b = caixa(sl, CX, y, w, hh, PAPEL2, raio=0.06)
    tf = b.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for lado_m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, lado_m, e(24))
    for i, l in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = l or ' '
        r.font.name = MONO; r.font.size = Pt(fs * 2 / 3)
        r.font.color.rgb = cor('7A7A72' if l.strip().startswith('#') else TINTA)
    b.name = 'Código: ' + d['alt'][:80]
    if lado:
        paragrafos(sl, CX + w + 44, y, CW - w - 44, CY + CH - y, lado, 28, ancora=MSO_ANCHOR.MIDDLE)


def s_fig(sl, a, d, ctx):
    y = titulo(sl, d['t'])
    ps = d.get('ps', [])
    png = ctx['png'][id(d)]
    if d.get('lado', True) and ps:
        w = CW * (1.25 / 2.25) - 22
        pic = figura(sl, png, CX, y, w, CY + CH - y)
        paragrafos(sl, CX + w + 44, y, CW - w - 44, CY + CH - y, ps, 28, ancora=MSO_ANCHOR.MIDDLE)
    else:
        hps = 0 if not ps else min(150, 50 * len(ps) + 40)
        pic = figura(sl, png, CX, y, CW, CY + CH - y - hps)
        if ps:
            paragrafos(sl, CX, CY + CH - hps, CW, hps, ps, 26)
    pic.name = 'Figura'
    pic._element.nvPicPr.cNvPr.set('descr', d['alt'])


def s_dois(sl, a, d, ctx):
    y = titulo(sl, d['t'])
    (l1, t1), (l2, t2) = d['lados']
    rod = d.get('rod')
    h = CY + CH - y - (80 if rod else 0)
    w = (CW - 22) / 2
    for i, (l, t, c) in enumerate(((l1, t1, d.get('c1', '')), (l2, t2, d.get('c2', 'y')))):
        b = caixa(sl, CX + i * (w + 22), y, w, h, CORES[c], raio=0.1)
        k = caber([l, t], w - 48, h - 40, 26, 14)
        tf = escrever(b, [f'**{l}**', t], k, ancora=MSO_ANCHOR.MIDDLE, margem=24, espaco=8)
        for r in tf.paragraphs[0].runs:
            r.font.size = Pt(k * 28 / 26 * 2 / 3)
    if rod:
        texto(sl, CX, CY + CH - 66, CW, 66, [rod], 26, alinhar=PP_ALIGN.CENTER)


def cronometro(sl, x, y, minutos):
    b = caixa(sl, x, y, 150, 42, TINTA, borda=False, raio=0.5)
    escrever(b, [f'{minutos} min'], 22, fonte=TIT, cor_txt=PAPEL, alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)


def s_ativ(sl, a, d, ctx):
    y = titulo(sl, 'Atividade: [[' + d['nome'] + ']]', w=CW - 180)
    cronometro(sl, CX + CW - 150, CY + 4, d['min'])
    lista(sl, y, d['passos'], marcador='num')


def s_check(sl, a, d, ctx):
    y = titulo(sl, '[[Confira]]: responda rápido', w=CW - 180)
    cronometro(sl, CX + CW - 150, CY + 4, 4)
    lista(sl, y, [q for q, _ in a['checagem']], marcador='num')


def s_levar(sl, a, d, ctx):
    y = titulo(sl, 'Para [[levar]]')
    lista(sl, y, a['levar'], rod=a.get('prox'))


def s_roteiro(sl, a, d, ctx):
    y = titulo(sl, 'Roteiro do [[professor]] · 30 min', 40)
    rot = a['roteiro']
    resp = ' '.join(f'**{i + 1}.** {r}' for i, (_, r) in enumerate(a['checagem']))
    hresp = 110
    tab = sl.shapes.add_table(len(rot), 3, e(CX), e(y), e(CW), e(CY + CH - y - hresp - 14)).table
    tab.first_row = False
    for j, w in enumerate((140, 260, CW - 400)):
        tab.columns[j].width = e(w)
    fs = caber([o for _, _, o in rot], CW - 420, CY + CH - y - hresp - 14, 21, 12, alt=1.32, extra=0.9 * len(rot))
    for i, (t, f, o) in enumerate(rot):
        for j, v in enumerate((f'**{t}**', f'**{f}**', o)):
            c = tab.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = cor(PAPEL if i % 2 else PAPEL2)
            c.margin_top = c.margin_bottom = e(5)
            c.text_frame.paragraphs[0].text = ''
            escrever(c, [v], fs, espaco=0)
    b = caixa(sl, CX, CY + CH - hresp, CW, hresp, PAPEL2, raio=0.12, espessura=2.5)
    k = caber(['Respostas do “Confira”: ' + resp], CW - 40, hresp - 20, 19, 11, alt=1.38)
    escrever(b, ['**Respostas do “Confira”:** ' + resp], k, ancora=MSO_ANCHOR.MIDDLE, margem=18)


TIPOS = {k[2:]: v for k, v in dict(globals()).items() if k.startswith('s_')}


def sequencia(a):
    seq = [('capa', {})] + a['slides'] + [('check', {}), ('levar', {}), ('roteiro', {})]
    if not any(t == 'hab' for t, _ in seq):
        i = 2 if len(seq) > 1 and seq[1][0] == 'perg' else 1
        seq.insert(i, ('hab', {}))
    tipos = [t for t, _ in seq]
    ini, fim = tipos.index('hab') + 2, tipos.index('ativ')
    a['roteiro'] = [(t, f, re.sub(r'\(slides \d+ a \d+\)', f'(slides {ini} a {fim})', o)) for t, f, o in a['roteiro']]
    return seq


def desenhar_svgs(aulas, fontes, pasta):
    """Converte todos os SVG das figuras em PNG (2x) com as fontes do site."""
    itens, mapa = [], {}
    for a in aulas:
        for t, d in a['slides']:
            if t == 'fig':
                nome = os.path.join(pasta, f'{a["cod"]}-{len(itens)}.png')
                itens.append({'svg': d['svg'], 'png': nome}); mapa[id(d)] = nome
    face = ''.join(f"@font-face{{font-family:'{fam}';font-weight:{w};src:url('file://{os.path.join(fontes, arq)}')}}"
                   for fam, w, arq in (('Lexend', 400, 'Lexend-400.ttf'), ('Lexend', 600, 'Lexend-600.ttf'),
                                       ('Sniglet', 800, 'Sniglet-800.ttf'), ('JetBrains Mono', 400, 'JetBrainsMono-Regular.ttf')))
    json.dump({'css': face, 'itens': itens}, open(os.path.join(pasta, 'svgs.json'), 'w'))
    js = r"""
const { chromium } = require('playwright'); const fs = require('fs');
(async () => {
  const { css, itens } = JSON.parse(fs.readFileSync(process.argv[2], 'utf8'));
  const b = await chromium.launch(); const p = await b.newPage({ deviceScaleFactor: 2 });
  for (const it of itens) {
    const arq = it.png.replace(/\.png$/, '.html');
    fs.writeFileSync(arq, `<html><head><style>${css} body{margin:0;background:transparent} svg{display:block}</style></head><body>${it.svg}</body></html>`);
    await p.goto('file://' + arq);
    await p.evaluate(() => document.fonts.ready);
    await (await p.$('svg')).screenshot({ path: it.png, omitBackground: true });
  }
  await b.close();
})();
"""
    open(os.path.join(pasta, 'svg.js'), 'w').write(js)
    subprocess.run(['node', os.path.join(pasta, 'svg.js'), os.path.join(pasta, 'svgs.json')], check=True)
    return mapa


def gerar(a, ctx, destino):
    prs = Presentation()
    prs.slide_width, prs.slide_height = e(1440), e(810)
    prs.core_properties.author = AUTORA
    prs.core_properties.title = f'BNCC Computação {a["cod"]}: {limpo(a["titulo"])}'
    seq = sequencia(a)
    for n, (t, d) in enumerate(seq, 1):
        sl = molde(prs, a, n, len(seq), t)
        TIPOS[t](sl, a, d, ctx)
        if t == 'roteiro' or n == 1:
            sl.notes_slide.notes_text_frame.text = f'{a["cod"]}: {a["hab"]}\nMaterial de {AUTORA} (ohtopico.com.br). Versão editável para apoiadores do Catarse.'
    nome = os.path.join(destino, f'bncc-{a["cod"].lower()}-slides.pptx')
    prs.save(nome)
    return nome


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--gerador', required=True); ap.add_argument('--fontes', required=True)
    ap.add_argument('--saida', default='_privado/premium/arquivos/slides'); ap.add_argument('--so', nargs='*')
    a = ap.parse_args()
    sys.path.insert(0, os.path.abspath(a.gerador))
    aulas = [x for m in ('aulas6', 'aulas7', 'aulas8', 'aulas9') for x in importlib.import_module(m).AULAS]
    if a.so:
        aulas = [x for x in aulas if x['cod'] in a.so]
    os.makedirs(a.saida, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        ctx = {'png': desenhar_svgs(aulas, os.path.abspath(a.fontes), tmp)}
        for x in aulas:
            print('ok', gerar(x, ctx, a.saida))
    arq = os.path.join(a.saida, 'titulos.json')  # títulos dos botões da área premium (proteger.py)
    titulos = json.load(open(arq, encoding='utf-8')) if os.path.exists(arq) else {}
    titulos.update({f'bncc-{x["cod"].lower()}-slides.pptx': limpo(x['titulo']) for x in aulas})
    json.dump(titulos, open(arq, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, sort_keys=True)


if __name__ == '__main__':
    main()
