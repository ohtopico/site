"""Kits mensais da área premium: para cada mês, 4 PDFs, 6 arquivos editáveis e um painel da aula em HTML.

O conteúdo vem de dois lugares no seu computador (fora do repositório, porque é material dos apoiadores):
  - o guia do professor de cada jogo (_privado/premium/meses/AAAA-MM/*.html e treine-a-maquina.html), lido por kits/extrair.py;
  - kits/dados.py, com objetivos, versão desplugada, ficha do aluno e avaliação de cada mês.
Saída, em _privado/premium/arquivos/AAAA-MM/:
  pdf/       desplugado, ficha do aluno, cartaz e avaliação com gabarito (plano Educador na prática)
  editaveis/ slides (.pptx), plano de aula, desplugado, ficha, avaliação e cartaz (.docx) (plano Educador criativo)
  html/      painel da aula para projetar (plano Educador maker)
e o trecho "Kits do mês" do hub.html, entre <!-- KITS --> e <!-- /KITS -->.

Uso (na raiz do repositório; precisa do Node com Playwright, python-docx e python-pptx):
    python3 _fonte/premium/editaveis/kit_mes.py --fontes _privado/premium --tipos CAMINHO/fontes [--meses 2026-12,2027-01]
--tipos é a pasta com Lexend-400.ttf, Lexend-600.ttf e Sniglet-800.ttf.
Autora: Júlia Roberta Quoos Alves.
"""
import argparse, html, json, os, re, subprocess, sys, tempfile
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Inches, Pt as PPt

AUTORA = 'Júlia Roberta Quoos Alves'
NOMES_MES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
CORES = ['#79c257', '#fff36b', '#f4b63b', '#e8503a', '#79c257']  # as cores do site, nunca azul
TINTA = '#111111'
PLANOS = {2: 'Educador na prática', 3: 'Educador criativo', 4: 'Educador maker'}
e = html.escape


def nome_mes(m):
    return f'{NOMES_MES[int(m[5:]) - 1]} de {m[:4]}'


def montar(mes, k, fontes):
    sys.path.insert(0, str(fontes / 'kits'))
    from extrair import jogo
    j = jogo(fontes / k['fonte'])
    t = (fontes / k['fonte']).read_text(encoding='utf-8')
    fases = [f for f in re.findall(r"nome: '([^']+)'", t) if f != 'Resultado']
    perg = next((re.sub(r'^Pergunta inicial: ', '', b) for _, b in j['roteiro'] if b.startswith('Pergunta inicial')), j['perguntas'][0])
    aspas = re.search(r'["“](.+?)["”]', perg)
    perg = k.get('pergunta') or (aspas.group(1) if aspas else re.sub(r'\.\s*Anote.*$', '', perg).strip())
    ideias = k.get('ideias') or j['ideias']
    j = {c: v for c, v in j.items() if c not in ('fases', 'ideias')}
    return dict(k, **j, mes=mes, nome_mes=nome_mes(mes), fases=fases, pergunta=perg, ideias=ideias, saida=re.sub(r'^Saída: ', '', j['roteiro'][-1][1]).capitalize()[:1] + re.sub(r'^Saída: ', '', j['roteiro'][-1][1])[1:])


# ---------------------------------------------------------------- PDF (HTML impresso pelo Chromium)
CSS = '''
@page { size: A4; margin: 14mm 15mm 20mm; }
@font-face { font-family: 'Lexend'; font-weight: 400; src: url('file://TIPOS/Lexend-400.ttf'); }
@font-face { font-family: 'Lexend'; font-weight: 600; src: url('file://TIPOS/Lexend-600.ttf'); }
@font-face { font-family: 'Sniglet'; font-weight: 800; src: url('file://TIPOS/Sniglet-800.ttf'); }
* { box-sizing: border-box; }
body { font-family: 'Lexend', sans-serif; font-size: 11.5pt; line-height: 1.45; color: #111; margin: 0; }
h1, h2, .big { font-family: 'Sniglet', 'Lexend', sans-serif; font-weight: 800; line-height: 1.1; }
h1 { font-size: 26pt; margin: 6px 0 4px; } h2 { font-size: 15pt; margin: 16px 0 6px; }
.faixa { display: flex; justify-content: space-between; border-bottom: 2px solid #111; padding-bottom: 6px; font-size: 9pt; }
.faixa b { background: #fff36b; border: 1.5px solid #111; border-radius: 999px; padding: 1px 10px; }
.sub { color: #444; margin: 0 0 10px; }
.caixas { display: flex; gap: 8px; flex-wrap: wrap; margin: 8px 0; }
.caixa { border: 1.5px solid #111; border-radius: 14px; padding: 6px 12px; font-size: 10pt; background: #f4f4f2; }
ol.passos { counter-reset: p; list-style: none; padding: 0; }
ol.passos li { counter-increment: p; position: relative; padding: 6px 0 6px 44px; min-height: 36px; break-inside: avoid; }
ol.passos li::before { content: counter(p); position: absolute; left: 0; top: 4px; width: 32px; height: 32px; border: 2px solid #111; border-radius: 50%;
  background: #79c257; font-family: 'Sniglet'; font-size: 15pt; display: flex; align-items: center; justify-content: center; }
.nome { display: grid; grid-template-columns: 3fr 1fr 1fr; gap: 12px; margin: 10px 0 4px; font-size: 10pt; }
.nome span { border-bottom: 1.5px solid #111; padding-bottom: 14px; }
.q { break-inside: avoid; margin-top: 12px; }
.q p { margin: 0 0 4px; font-weight: 600; }
.linha { border-bottom: 1px solid #888; height: 26px; }
.alts { list-style: none; padding: 0; margin: 4px 0 0; }
.alts li { padding: 2px 0 2px 30px; position: relative; }
.alts li span { position: absolute; left: 0; width: 22px; height: 22px; border: 1.5px solid #111; border-radius: 50%; font-size: 9pt; display: flex; align-items: center; justify-content: center; }
.leve { border: 2px solid #111; border-radius: 16px; padding: 10px 14px; margin-top: 16px; background: #fffbd6; break-inside: avoid; }
table { border-collapse: collapse; width: 100%; font-size: 10pt; }
th, td { border: 1.5px solid #111; padding: 6px 8px; text-align: left; vertical-align: top; }
th { background: #fff36b; }
.quebra { break-before: page; }
.cartaz { height: 236mm; display: flex; flex-direction: column; }
.cartaz h1 { font-size: 44pt; margin: 10px 0 2px; }
.cartaz .ideia { flex: 1; display: flex; gap: 14px; align-items: center; border: 2.5px solid #111; border-radius: 22px; padding: 10px 18px; margin-top: 10px; font-size: 15pt; line-height: 1.3; }
.cartaz .ideia .n { font-family: 'Sniglet'; font-size: 36pt; min-width: 44px; }
'''


def doc_html(titulo, corpo, d):
    return (f'<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><title>{e(titulo)}</title><meta name="author" content="{AUTORA}"><style>{CSS}</style></head><body>'
            f'<div class="faixa"><span>ohtopico · Kit de {d["nome_mes"]}</span><b>{e(d["titulo"])}</b></div>{corpo}'
            '</body></html>')


def bncc_txt(d):
    return ', '.join(c for c, _ in d['bncc'])


def pdf_desplugado(d):
    p = d['desplugado']
    corpo = (f'<h1>{e(p["titulo"])}</h1><p class="sub">Versão desplugada de {e(d["titulo"])}: a ideia do jogo sem computador nem internet.</p>'
             f'<div class="caixas"><span class="caixa">Tempo: {e(p["tempo"])}</span><span class="caixa">BNCC: {bncc_txt(d)}</span><span class="caixa">Em grupos de 3 a 5</span></div>'
             f'<p>{e(d["etapa"])}</p><h2>Materiais</h2><ul>{"".join(f"<li>{e(x)}</li>" for x in p["materiais"])}</ul>'
             f'<h2>Antes da aula</h2><p>{e(p["preparo"])}</p><h2>Passo a passo</h2><ol class="passos">{"".join(f"<li>{e(x)}</li>" for x in p["passos"])}</ol>'
             f'<h2>Fechamento</h2><p>{e(p["fechamento"])}</p><h2>Perguntas para o debate</h2><ol>{"".join(f"<li>{e(x)}</li>" for x in d["perguntas"])}</ol>')
    return doc_html(p['titulo'], corpo, d)


def pdf_ficha(d):
    qs = ''.join(f'<div class="q"><p>{i}. {e(q)}</p>{"<div class=linha></div>" * n}</div>' for i, (q, n) in enumerate(d['ficha'], 1))
    corpo = (f'<div class="nome"><span>Nome:</span><span>Turma:</span><span>Data:</span></div><h1>Ficha do aluno</h1>'
             f'<p class="sub">Responda enquanto joga {e(d["titulo"])} ou logo depois.</p>{qs}'
             f'<div class="leve"><b>O que eu levo desta aula</b>{"<div class=linha></div>" * 2}</div>')
    return doc_html('Ficha do aluno', corpo, d)


def pdf_cartaz(d):
    ideias = ''.join(f'<div class="ideia" style="background:{CORES[i % 5]};{"color:#111" if i % 5 != 3 else "color:#fff"}"><span class="n">{i + 1}</span><span>{ideia_html(x)}</span></div>'
                     for i, x in enumerate(d['ideias'][:5]))
    corpo = f'<div class="cartaz"><h1>{e(d["titulo"])}</h1><p class="sub" style="font-size:13pt">{"O que aprendemos" if len(d["ideias"]) else ""}</p>{ideias}</div>'
    return doc_html('Cartaz', corpo, d)


def ideia_html(x):
    """Primeira frase em negrito, como no resultado dos jogos."""
    m = re.match(r'^(.+?[.:])\s+(.*)$', x)
    return f'<b>{e(m.group(1))}</b> {e(m.group(2))}' if m else e(x)


def pdf_avaliacao(d):
    letras = 'abcd'
    objs = ''.join(f'<div class="q"><p>{i}. {e(q)}</p><ul class="alts">{"".join(f"<li><span>{letras[a]}</span>{e(x)}</li>" for a, x in enumerate(alts))}</ul></div>'
                   for i, (q, alts, _) in enumerate(d['objetivas'], 1))
    q6, rub = d['discursiva']
    n = len(d['objetivas']) + 1
    gab = ''.join(f'<li>Questão {i}: <b>{letras[c]}</b> ({e(alts[c])})</li>' for i, (_, alts, c) in enumerate(d['objetivas'], 1))
    corpo = (f'<div class="nome"><span>Nome:</span><span>Turma:</span><span>Data:</span></div><h1>Avaliação</h1>'
             f'<p class="sub">{e(d["titulo"])} · BNCC {bncc_txt(d)}. Marque uma alternativa nas questões 1 a {n - 1}.</p>{objs}'
             f'<div class="q"><p>{n}. {e(q6)}</p>{"<div class=linha></div>" * 7}</div>'
             f'<div class="quebra"></div><h1 style="font-size:20pt">Gabarito e rubrica</h1><p class="sub">Para o professor. Sugestão de pontuação: 1 ponto por questão objetiva e até 5 pontos na questão {n}.</p>'
             f'<ul>{gab}</ul><h2>Questão {n}</h2><table><tr><th style="width:22%">Nível</th><th>O que a resposta mostra</th></tr>'
             f'<tr><td>Completa (5 pontos)</td><td>{e(rub[0])}</td></tr><tr><td>Parcial (3 pontos)</td><td>{e(rub[1])}</td></tr><tr><td>Inicial (1 ponto)</td><td>{e(rub[2])}</td></tr></table>'
             f'<h2>Respostas esperadas do jogo</h2><ul>{"".join(f"<li>{e(x)}</li>" for x in d["respostas"])}</ul>')
    return doc_html('Avaliação', corpo, d)


def imprimir(pares, tipos):
    with tempfile.TemporaryDirectory() as tmp:
        lista = []
        for i, (conteudo, destino) in enumerate(pares):
            f = os.path.join(tmp, f'{i}.html'); open(f, 'w', encoding='utf-8').write(conteudo.replace('TIPOS', str(Path(tipos).resolve())))
            lista.append([f, str(Path(destino).resolve())])
        rod = (f'<div style="font-family:sans-serif;font-size:7.5px;color:#555;width:100%;padding:0 15mm;display:flex;justify-content:space-between">'
               f'<span>Material de {AUTORA} para os apoiadores do Catarse (ohtopico.com.br). Uso livre em sala de aula, com citação da fonte. Por favor, não compartilhe os arquivos.</span>'
               '<span><span class="pageNumber"></span>/<span class="totalPages"></span></span></div>')
        js = (f"const RODAPE={json.dumps(rod)};" "const {chromium}=require('playwright');(async()=>{const b=await chromium.launch();const p=await b.newPage();"
              f"for (const [h,o] of {json.dumps(lista)}) {{await p.goto('file://'+h);await p.waitForTimeout(150);"
              "await p.pdf({path:o,format:'A4',printBackground:true,tagged:true,displayHeaderFooter:true,headerTemplate:'<span></span>',footerTemplate:RODAPE});}await b.close();})();")
        open(os.path.join(tmp, 'pdf.js'), 'w').write(js)
        subprocess.run(['node', os.path.join(tmp, 'pdf.js')], check=True)


# ---------------------------------------------------------------- Word
def word(d):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    for nome, tam in (('Title', 24), ('Heading 1', 16), ('Heading 2', 13)):
        s = doc.styles[nome]; s.font.name = 'Arial'; s.font.size = Pt(tam); s.font.bold = True; s.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        rpr = s.element.get_or_add_rPr(); f = rpr.find(qn('w:rFonts'))
        if f is None:
            f = OxmlElement('w:rFonts'); rpr.append(f)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            f.set(qn(a), 'Arial')
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if f.get(qn(a)) is not None:
                del f.attrib[qn(a)]
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.8); sec.left_margin = sec.right_margin = Cm(2)
    doc.core_properties.author = AUTORA
    return doc


def rodape(doc):
    p = doc.add_paragraph(); r = p.add_run(f'Material de {AUTORA} para os apoiadores do Catarse (ohtopico.com.br). Uso livre em sala de aula, com citação da fonte.')
    r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def linhas(doc, n):
    for _ in range(n):
        doc.add_paragraph('_' * 86)


def cabecalho(doc, d, titulo):
    doc.add_paragraph(f'ohtopico · Kit de {d["nome_mes"]} · {d["titulo"]}').runs[0].font.size = Pt(9)
    doc.add_heading(titulo, level=0)


def tabela(doc, linhas_, cabecalho_=True):
    t = doc.add_table(rows=0, cols=len(linhas_[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, l in enumerate(linhas_):
        cel = t.add_row().cells
        for c, v in zip(cel, l):
            c.text = v
            if cabecalho_ and i == 0:
                c.paragraphs[0].runs[0].bold = True
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'FFF36B'); c._tc.get_or_add_tcPr().append(sh)
    return t


def docx_plano(d, destino):
    doc = word(d); cabecalho(doc, d, f'Plano de aula: {d["titulo"]}')
    doc.add_paragraph(d['etapa'])
    doc.add_heading('Objetivos', 1)
    for x in d['objetivos']:
        doc.add_paragraph(x, style='List Bullet')
    doc.add_heading('Habilidades da BNCC', 1)
    for c, x in d['bncc']:
        p = doc.add_paragraph(style='List Bullet'); p.add_run(c + ' ').bold = True; p.add_run(x)
    doc.add_heading('Materiais', 1)
    for x in ['Um computador ou celular por dupla, com o jogo ' + d['titulo'] + ' (área premium ou arquivo .html baixado)'] + d['desplugado']['materiais']:
        doc.add_paragraph(x, style='List Bullet')
    doc.add_heading('Pergunta inicial', 1); doc.add_paragraph(d['pergunta'])
    doc.add_heading('Roteiro de 2 aulas', 1); tabela(doc, [['Tempo', 'O que fazer']] + [list(x) for x in d['roteiro']])
    doc.add_heading('As fases do jogo', 1)
    for i, f in enumerate(d['fases'], 1):
        doc.add_paragraph(f'Fase {i}: {f}', style='List Number')
    for t, x in d['secoes']:
        doc.add_heading(t, 1); doc.add_paragraph(x)
    doc.add_heading('Perguntas para o debate', 1)
    for x in d['perguntas']:
        doc.add_paragraph(x, style='List Number')
    doc.add_heading('Respostas esperadas', 1)
    for x in d['respostas']:
        doc.add_paragraph(x, style='List Bullet')
    doc.add_heading('Avaliação', 1)
    doc.add_paragraph('Use a avaliação do kit (5 questões objetivas e 1 discursiva com rubrica) e a ficha do aluno preenchida durante o jogo. '
                      'Observe também a participação no debate e a regra escrita na atividade de saída.')
    rodape(doc); doc.save(destino)


def docx_desplugado(d, destino):
    p = d['desplugado']; doc = word(d); cabecalho(doc, d, p['titulo'])
    doc.add_paragraph(f'Versão desplugada de {d["titulo"]}. Tempo: {p["tempo"]}. BNCC: {bncc_txt(d)}.')
    doc.add_heading('Materiais', 1)
    for x in p['materiais']:
        doc.add_paragraph(x, style='List Bullet')
    doc.add_heading('Antes da aula', 1); doc.add_paragraph(p['preparo'])
    doc.add_heading('Passo a passo', 1)
    for x in p['passos']:
        doc.add_paragraph(x, style='List Number')
    doc.add_heading('Fechamento', 1); doc.add_paragraph(p['fechamento'])
    rodape(doc); doc.save(destino)


def docx_ficha(d, destino):
    doc = word(d); doc.add_paragraph('Nome: ______________________________________   Turma: ________   Data: ___/___/_____')
    cabecalho(doc, d, 'Ficha do aluno')
    for i, (q, n) in enumerate(d['ficha'], 1):
        doc.add_paragraph().add_run(f'{i}. {q}').bold = True; linhas(doc, n)
    doc.add_paragraph().add_run('O que eu levo desta aula').bold = True; linhas(doc, 2)
    rodape(doc); doc.save(destino)


def docx_avaliacao(d, destino):
    doc = word(d); doc.add_paragraph('Nome: ______________________________________   Turma: ________   Data: ___/___/_____')
    cabecalho(doc, d, 'Avaliação')
    for i, (q, alts, _) in enumerate(d['objetivas'], 1):
        doc.add_paragraph().add_run(f'{i}. {q}').bold = True
        for a, x in zip('abcd', alts):
            doc.add_paragraph(f'( {a} ) {x}')
    n = len(d['objetivas']) + 1
    doc.add_paragraph().add_run(f'{n}. {d["discursiva"][0]}').bold = True; linhas(doc, 7)
    doc.add_page_break(); doc.add_heading('Gabarito e rubrica (para o professor)', 1)
    for i, (_, alts, c) in enumerate(d['objetivas'], 1):
        doc.add_paragraph(f'Questão {i}: {"abcd"[c]} ({alts[c]})', style='List Bullet')
    r = d['discursiva'][1]
    tabela(doc, [['Nível', f'Questão {n}'], ['Completa (5 pontos)', r[0]], ['Parcial (3 pontos)', r[1]], ['Inicial (1 ponto)', r[2]]])
    rodape(doc); doc.save(destino)


def docx_cartaz(d, destino):
    doc = word(d); h = doc.add_heading(d['titulo'], level=0)
    for i, x in enumerate(d['ideias'][:5], 1):
        p = doc.add_paragraph(); r = p.add_run(f'{i}. {x}'); r.font.size = Pt(20)
        p.paragraph_format.space_after = Pt(18)
    rodape(doc); doc.save(destino)


# ---------------------------------------------------------------- PowerPoint
def pptx_slides(d, destino):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    branco = prs.slide_layouts[6]
    hexrgb = lambda h: PRGB.from_string(h.lstrip('#').upper())

    def slide(fundo='#ffffff'):
        s = prs.slides.add_slide(branco); s.background.fill.solid(); s.background.fill.fore_color.rgb = hexrgb(fundo); return s

    def texto(s, x, y, w, h, txt, tam=28, bold=False, cor=TINTA):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        for i, linha in enumerate(txt if isinstance(txt, list) else [txt]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = linha; r.font.size = PPt(tam); r.font.bold = bold; r.font.name = 'Arial'; r.font.color.rgb = hexrgb(cor)
            p.space_after = PPt(tam * 0.5)
        return tb

    def titulo(s, t):
        texto(s, 0.7, 0.45, 12, 1.1, t, 40, True)

    s = slide('#fff36b'); texto(s, 0.8, 2.0, 11.5, 1.6, d['titulo'], 60, True)
    texto(s, 0.8, 3.7, 11.5, 1.6, d['lead'], 20); texto(s, 0.8, 6.4, 11.5, 0.6, f'ohtopico · {d["nome_mes"]} · {AUTORA}', 14)
    s = slide('#79c257'); texto(s, 0.8, 0.6, 11.5, 0.8, 'Para começar', 24, True); texto(s, 0.8, 2.0, 11.7, 3.5, d['pergunta'], 48, True)
    s = slide(); titulo(s, 'Objetivos da aula'); texto(s, 0.8, 1.8, 11.7, 5, ['• ' + x for x in d['objetivos']], 28)
    s = slide(); titulo(s, 'As fases do jogo'); texto(s, 0.8, 1.8, 11.7, 5, [f'{i}. {f}' for i, f in enumerate(d['fases'], 1)], 32)
    for i, x in enumerate(d['ideias'][:5]):
        c = CORES[i % 5]; s = slide(c); m = re.match(r'^(.+?[.:])\s+(.*)$', x)
        cor = '#ffffff' if i % 5 == 3 else TINTA
        texto(s, 0.8, 0.6, 11.5, 0.8, f'O que aprendemos · {i + 1}', 22, True, cor)
        texto(s, 0.8, 2.0, 11.7, 1.8, m.group(1) if m else x, 46, True, cor)
        if m:
            texto(s, 0.8, 4.0, 11.7, 2.5, m.group(2), 26, False, cor)
    p = d['desplugado']; s = slide(); titulo(s, f'Sem computador: {p["titulo"]}')
    texto(s, 0.8, 1.7, 11.7, 5.5, [f'{i}. {x}' for i, x in enumerate(p['passos'], 1)], 22)
    pergs = d['perguntas']
    passo = len(pergs) if len(pergs) <= 3 else -(-len(pergs) // 2)
    for parte in range(0, len(pergs), passo):
        s = slide(); titulo(s, 'Para debater'); texto(s, 0.8, 1.8, 11.7, 5.3, [f'{parte + i + 1}. {x}' for i, x in enumerate(pergs[parte:parte + passo])], 26)
    s = slide('#f4b63b'); texto(s, 0.8, 0.6, 11.5, 0.8, 'Para sair', 24, True); texto(s, 0.8, 2.0, 11.7, 4, d['saida'], 40, True)
    prs.core_properties.author = AUTORA
    prs.save(destino)


# ---------------------------------------------------------------- painel da aula (HTML para projetar)
def painel(d):
    dados = json.dumps(dict(titulo=d['titulo'], mes=d['nome_mes'], pergunta=d['pergunta'], fases=d['fases'], ideias=d['ideias'][:5],
                            perguntas=d['perguntas'], passos=d['desplugado']['passos'], saida=d['saida'], slug=d['slug']), ensure_ascii=False)
    return '''<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<meta name="robots" content="noindex, nofollow">
<meta name="author" content="''' + AUTORA + '''">
<title>Painel da aula · ''' + e(d['titulo']) + '''</title>
<style>
  :root { --ink: #111; --paper: #fffdf6; --verde: #79c257; --amarelo: #fff36b; --laranja: #f4b63b; --vermelho: #e8503a; }
  * { box-sizing: border-box; }
  body { margin: 0; font-family: system-ui, "Segoe UI", Arial, sans-serif; background: var(--paper); color: var(--ink); }
  header { display: flex; flex-wrap: wrap; gap: 8px; align-items: center; justify-content: space-between; padding: 14px 18px; border-bottom: 2px solid var(--ink); }
  header h1 { margin: 0; font-size: clamp(20px, 3vw, 30px); }
  nav { display: flex; flex-wrap: wrap; gap: 6px; }
  button { font: inherit; border: 2px solid var(--ink); border-radius: 999px; padding: 8px 16px; background: #fff; color: var(--ink); cursor: pointer; }
  button[aria-pressed="true"], button.cheio { background: var(--ink); color: #fff; }
  main { padding: 20px 18px 40px; max-width: 1100px; margin: 0 auto; }
  .tela { display: none; } .tela.ativa { display: block; }
  .grande { font-size: clamp(28px, 5vw, 60px); font-weight: 800; line-height: 1.15; margin: 24px 0; }
  .cartao { border: 3px solid var(--ink); border-radius: 26px; padding: 22px 26px; margin: 14px 0; box-shadow: 8px 8px 0 var(--ink); font-size: clamp(18px, 2.6vw, 30px); }
  .relogio { font-size: clamp(80px, 18vw, 220px); font-weight: 800; text-align: center; font-variant-numeric: tabular-nums; margin: 10px 0; }
  .relogio.fim { color: var(--vermelho); }
  .linha { display: flex; flex-wrap: wrap; gap: 8px; justify-content: center; }
  textarea { width: 100%; min-height: 120px; font: inherit; font-size: 18px; border: 2px solid var(--ink); border-radius: 16px; padding: 10px; }
  .duplas { display: grid; grid-template-columns: repeat(auto-fill, minmax(220px, 1fr)); gap: 10px; margin-top: 14px; }
  .duplas div { border: 2px solid var(--ink); border-radius: 16px; padding: 10px 14px; background: var(--amarelo); font-size: 20px; }
  .placar { display: grid; grid-template-columns: repeat(auto-fill, minmax(200px, 1fr)); gap: 10px; }
  .placar div { border: 2px solid var(--ink); border-radius: 16px; padding: 10px; text-align: center; font-size: 20px; }
  .placar b { display: block; font-size: 40px; }
  .muted { color: #555; }
  @media print { header, .linha { display: none; } .tela { display: block; } }
</style>
</head>
<body>
<header><h1 id="t"></h1><nav aria-label="Partes da aula" id="nav"></nav></header>
<main>
  <section class="tela" id="s-pergunta"><p class="muted">Para começar</p><p class="grande" id="pergunta"></p></section>
  <section class="tela" id="s-tempo"><p class="muted">Cronômetro</p><div class="relogio" id="relogio" role="timer" aria-live="off">10:00</div>
    <div class="linha"><button data-min="5">5 min</button><button data-min="10">10 min</button><button data-min="15">15 min</button><button data-min="25">25 min</button><button id="pausa">Pausar</button></div></section>
  <section class="tela" id="s-duplas"><p class="muted">Cole os nomes da turma, um por linha, e sorteie as duplas. Os nomes ficam só nesta página.</p>
    <label for="nomes" class="muted">Nomes</label><textarea id="nomes"></textarea><div class="linha" style="margin-top:10px"><button id="sortear" class="cheio">Sortear duplas</button></div><div class="duplas" id="duplas"></div></section>
  <section class="tela" id="s-placar"><p class="muted">Estrelas de cada dupla nas fases do jogo. Toque para somar uma estrela.</p><div class="placar" id="placar"></div></section>
  <section class="tela" id="s-ideias"><p class="muted">O que aprendemos</p><div id="ideias"></div><div class="linha"><button id="mais" class="cheio">Mostrar a próxima</button></div></section>
  <section class="tela" id="s-debate"><p class="muted">Para debater</p><p class="grande" id="debate"></p><div class="linha"><button id="ant">Anterior</button><button id="prox" class="cheio">Próxima pergunta</button></div></section>
  <section class="tela" id="s-saida"><p class="muted">Para sair</p><p class="grande" id="saida"></p></section>
  <p class="muted" style="margin-top:40px;font-size:13px">Painel criado por ''' + AUTORA + ''' para os apoiadores do Catarse (ohtopico.com.br). Funciona sem internet. Uso livre em sala de aula, com citação da fonte.</p>
</main>
<script>
(function () {
  const D = ''' + dados + ''';
  const $ = id => document.getElementById(id);
  $('t').textContent = D.titulo; document.title = 'Painel da aula · ' + D.titulo;
  $('pergunta').textContent = D.pergunta; $('saida').textContent = D.saida;
  const TELAS = [['pergunta', 'Pergunta'], ['tempo', 'Cronômetro'], ['duplas', 'Duplas'], ['placar', 'Placar'], ['ideias', 'Ideias'], ['debate', 'Debate'], ['saida', 'Saída']];
  $('nav').innerHTML = TELAS.map(([k, t]) => `<button type="button" data-tela="${k}">${t}</button>`).join('');
  function ver(k) { TELAS.forEach(([x]) => $('s-' + x).classList.toggle('ativa', x === k)); document.querySelectorAll('[data-tela]').forEach(b => b.setAttribute('aria-pressed', b.dataset.tela === k)); }
  document.querySelectorAll('[data-tela]').forEach(b => b.onclick = () => ver(b.dataset.tela)); ver('pergunta');
  let resta = 600, rodando = null;
  const mostra = () => { const r = $('relogio'); r.textContent = String(Math.floor(resta / 60)).padStart(2, '0') + ':' + String(resta % 60).padStart(2, '0'); r.classList.toggle('fim', resta === 0); };
  function iniciar(min) { resta = min * 60; clearInterval(rodando); rodando = setInterval(() => { if (resta > 0) { resta--; mostra(); } else clearInterval(rodando); }, 1000); mostra(); $('pausa').textContent = 'Pausar'; }
  document.querySelectorAll('[data-min]').forEach(b => b.onclick = () => iniciar(+b.dataset.min));
  $('pausa').onclick = () => { if (rodando) { clearInterval(rodando); rodando = null; $('pausa').textContent = 'Continuar'; } else iniciar(resta / 60); };
  let duplas = [];
  $('sortear').onclick = () => {
    const n = $('nomes').value.split('\\n').map(s => s.trim()).filter(Boolean);
    for (let i = n.length - 1; i > 0; i--) { const j = Math.floor(Math.random() * (i + 1)); [n[i], n[j]] = [n[j], n[i]]; }
    duplas = []; for (let i = 0; i < n.length; i += 2) duplas.push(n.slice(i, i + 2));
    if (duplas.length > 1 && duplas[duplas.length - 1].length === 1) duplas[duplas.length - 2].push(duplas.pop()[0]);
    $('duplas').innerHTML = duplas.map((d, i) => `<div><b>Dupla ${i + 1}</b><br>${d.map(x => x.replace(/[<>&]/g, '')).join(' e ')}</div>`).join('');
    placar();
  };
  const estrelas = {};
  function placar() {
    const lista = duplas.length ? duplas.map((_, i) => 'Dupla ' + (i + 1)) : Array.from({ length: 12 }, (_, i) => 'Dupla ' + (i + 1));
    $('placar').innerHTML = lista.map(n => `<button type="button" data-d="${n}"><b>${estrelas[n] || 0}</b>${n}</button>`).join('');
    document.querySelectorAll('[data-d]').forEach(b => b.onclick = () => { estrelas[b.dataset.d] = (estrelas[b.dataset.d] || 0) + 1; placar(); });
  }
  placar();
  let vis = 0; const cores = ['#79c257', '#fff36b', '#f4b63b', '#e8503a', '#79c257'];
  function ideias() { $('ideias').innerHTML = D.ideias.slice(0, vis).map((x, i) => `<div class="cartao" style="background:${cores[i % 5]};color:${i % 5 === 3 ? '#fff' : '#111'}">${x.replace(/[<>&]/g, '')}</div>`).join(''); $('mais').hidden = vis >= D.ideias.length; }
  $('mais').onclick = () => { vis++; ideias(); }; ideias();
  let q = 0; const debate = () => { $('debate').textContent = (q + 1) + '. ' + D.perguntas[q]; $('ant').disabled = q === 0; $('prox').disabled = q === D.perguntas.length - 1; };
  $('ant').onclick = () => { q--; debate(); }; $('prox').onclick = () => { q++; debate(); }; debate();
  document.addEventListener('keydown', ev => { if (ev.target.tagName === 'TEXTAREA') return; const i = TELAS.findIndex(([k]) => $('s-' + k).classList.contains('ativa'));
    if (ev.key === 'ArrowRight' && i < TELAS.length - 1) ver(TELAS[i + 1][0]); if (ev.key === 'ArrowLeft' && i > 0) ver(TELAS[i - 1][0]); });
})();
</script>
</body>
</html>
'''


# ---------------------------------------------------------------- hub
def hub_trecho(meses):
    blocos = []
    for d in meses:
        m, s = d['mes'], d['slug']
        pdfs = [('1-desplugado', 'Versão desplugada'), ('2-ficha-do-aluno', 'Ficha do aluno'), ('3-cartaz', 'Cartaz'), ('4-avaliacao', 'Avaliação com gabarito')]
        if d.get('sem_desplugado'):
            pdfs = pdfs[1:]
        eds = [('slides.pptx', 'Slides da aula (PowerPoint)'), ('plano-de-aula.docx', 'Plano de aula'), ('desplugado.docx', 'Versão desplugada'),
               ('ficha-do-aluno.docx', 'Ficha do aluno'), ('avaliacao.docx', 'Avaliação com gabarito'), ('cartaz.docx', 'Cartaz')]
        bot = lambda rel, t: f'<button class="baixar" type="button" data-arquivo="{rel}">{t} <small>{{{{TAMANHO:{rel}}}}}</small></button>'
        blocos.append(f'''    <details class="kit"{' open' if not blocos else ''}>
      <summary>{d["nome_mes"].capitalize()} · {e(d["titulo"])}</summary>
      <div class="grid">
        <div class="item" data-nivel="2" data-mes="{m}">
          <span class="tipo">PDF · para imprimir</span>
          <strong>Pacote de {len(pdfs)} PDFs</strong>
          <span>{'Ficha do aluno, cartaz para a sala e avaliação com gabarito e rubrica. A versão desplugada deste mês está em Guia do professor e materiais para imprimir.' if d.get('sem_desplugado') else f'Versão desplugada ({e(d["desplugado"]["titulo"])}), ficha do aluno, cartaz para a sala e avaliação com gabarito e rubrica.'}</span>
          <div class="acoes">{''.join(bot(f'{m}/pdf/{s}-{a}.pdf', t) for a, t in pdfs)}</div>
        </div>
        <div class="item" data-nivel="3" data-mes="{m}">
          <span class="tipo">PowerPoint e Word · editáveis</span>
          <strong>Pacote de 6 arquivos editáveis</strong>
          <span>Slides da aula, plano de aula com roteiro e BNCC, versão desplugada, ficha, avaliação e cartaz, para adaptar à sua turma.</span>
          <div class="acoes">{''.join(bot(f'{m}/editaveis/{s}-{a}', t) for a, t in eds)}</div>
        </div>
        <div class="item" data-nivel="4" data-mes="{m}">
          <span class="tipo">HTML · para projetar</span>
          <strong>Painel da aula</strong>
          <span>Página para projetar com pergunta inicial, cronômetro, sorteio de duplas, placar de estrelas, ideias do jogo e perguntas do debate. Funciona sem internet. É o segundo recurso em HTML do mês, junto com o jogo.</span>
          <div class="acoes">{bot(f'{m}/html/{s}-painel-da-aula.html', 'Baixar o painel (.html)')}</div>
        </div>
      </div>
    </details>''')
    return ('<!-- KITS -->\n  <h2 id="kits">Kits do mês</h2>\n  <p class="muted intro">Todo mês, um kit em volta do jogo do mês: PDFs para imprimir, arquivos editáveis e um painel da aula. '
            'Cada kit abre no dia 1 do seu mês, e cada parte abre para o seu plano e os de cima.</p>\n' + '\n'.join(blocos) + '\n  <!-- /KITS -->')


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--fontes', default='_privado/premium'); ap.add_argument('--tipos', required=True); ap.add_argument('--meses', default='')
    a = ap.parse_args(); fontes = Path(a.fontes)
    sys.path.insert(0, str(fontes / 'kits'))
    from dados import KITS
    todos = [montar(m, k, fontes) for m, k in sorted(KITS.items())]
    escolhidos = [d for d in todos if not a.meses or d['mes'] in a.meses.split(',')]
    pares = []
    for d in escolhidos:
        base = fontes / 'arquivos' / d['mes']; s = d['slug']
        for p in ('pdf', 'editaveis', 'html'):
            (base / p).mkdir(parents=True, exist_ok=True)
        if not d.get('sem_desplugado'):
            pares.append((pdf_desplugado(d), base / 'pdf' / f'{s}-1-desplugado.pdf'))
        pares += [(pdf_ficha(d), base / 'pdf' / f'{s}-2-ficha-do-aluno.pdf'), (pdf_cartaz(d), base / 'pdf' / f'{s}-3-cartaz.pdf'), (pdf_avaliacao(d), base / 'pdf' / f'{s}-4-avaliacao.pdf')]
        pptx_slides(d, base / 'editaveis' / f'{s}-slides.pptx')
        docx_plano(d, base / 'editaveis' / f'{s}-plano-de-aula.docx'); docx_desplugado(d, base / 'editaveis' / f'{s}-desplugado.docx')
        docx_ficha(d, base / 'editaveis' / f'{s}-ficha-do-aluno.docx'); docx_avaliacao(d, base / 'editaveis' / f'{s}-avaliacao.docx'); docx_cartaz(d, base / 'editaveis' / f'{s}-cartaz.docx')
        (base / 'html' / f'{s}-painel-da-aula.html').write_text(painel(d), encoding='utf-8')
        print('ok kit', d['mes'], s)
    imprimir(pares, a.tipos)
    niveis = fontes / 'arquivos' / 'niveis.json'; nv = json.loads(niveis.read_text())
    nv.update(editaveis=3, html=4); niveis.write_text(json.dumps(nv))
    hub = fontes / 'hub.html'; t = hub.read_text(encoding='utf-8'); trecho = hub_trecho(todos)
    t = re.sub(r'<!-- KITS -->.*?<!-- /KITS -->', lambda _: trecho, t, flags=re.S) if '<!-- KITS -->' in t else t.replace('  <h2>Calendário de jogos exclusivos</h2>', trecho + '\n\n  <h2>Calendário de jogos exclusivos</h2>', 1)
    hub.write_text(t, encoding='utf-8')
    print('ok hub e niveis.json')


if __name__ == '__main__':
    main()
