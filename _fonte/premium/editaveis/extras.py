"""Materiais extras da área premium, fora dos kits mensais:
  - Kit Feira de Ciências com IA: roteiro de 6 semanas, ideias de projeto, ficha de inscrição, diário de bordo, roteiro do banner,
    cuidados com dados e ética, rubrica, ficha do visitante e certificado;
  - Guia de adaptação dos kits: orientações por perfil (TEA, TDAH, dislexia, deficiência intelectual, baixa visão, cegueira,
    surdez e deficiência motora) e lista de conferência para planejar aulas acessíveis.
Saída, em _privado/premium/arquivos/: imprimir/ (PDF, plano Educador na prática) e editar/ (Word, plano Educador criativo).

Uso (na raiz do repositório; precisa do Node com Playwright e de python-docx):
    python3 _fonte/premium/editaveis/extras.py --fontes _privado/premium --tipos CAMINHO/fontes
Autora: Júlia Roberta Quoos Alves.
"""
import argparse, html, json, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from kit_mes import AUTORA, CSS, imprimir, word, rodape, linhas, tabela
from docx.shared import Pt

e = html.escape

# ---------------------------------------------------------------- conteúdo: Feira de Ciências com IA
FEIRA = dict(
    titulo='Feira de Ciências com IA',
    lead='Um roteiro de 6 semanas para a turma criar, testar e apresentar um projeto que usa inteligência artificial ou dados para responder a uma pergunta real da escola ou do bairro. Para o 6º ao 9º ano, em grupos de 3 a 5 estudantes.',
    bncc=[('EF07CO03', 'Construir soluções computacionais de problemas de diferentes áreas, individual e colaborativamente.'),
          ('EF08CO04', 'Construir soluções computacionais de problemas de diferentes áreas, individual e colaborativamente.'),
          ('EF09CO02', 'Construir soluções computacionais de problemas de diferentes áreas, individual e colaborativamente.'),
          ('EF09CO06', 'Analisar problemas sociais da cidade e do estado a partir de ambientes digitais, propondo soluções.'),
          ('EF09CO07', 'Avaliar aplicações e implicações políticas, socioambientais e culturais das tecnologias digitais.')],
    semanas=[
        ('1', 'Pergunta e grupos', 'Formar os grupos, escolher o tema e escrever a pergunta do projeto. Preencher a ficha de inscrição.', 'Ficha de inscrição pronta'),
        ('2', 'Dados e pesquisa', 'Descobrir o que já se sabe sobre o tema e decidir quais dados coletar (fotos, medidas, respostas, sons) e como coletar com cuidado.', 'Plano de coleta e primeiros dados'),
        ('3', 'Protótipo', 'Montar a primeira versão: treinar um modelo simples, montar uma planilha com gráfico ou criar regras de decisão. Pode ser no computador ou desplugado.', 'Primeira versão funcionando'),
        ('4', 'Testes', 'Testar com dados novos, contar acertos e erros, anotar o que melhorou de uma versão para a outra e o que ainda confunde o modelo.', 'Tabela de testes e versão melhorada'),
        ('5', 'Banner e ensaio', 'Montar o banner ou o painel digital, preparar a fala de 3 minutos e ensaiar com outro grupo fazendo perguntas.', 'Banner pronto e fala ensaiada'),
        ('6', 'Feira', 'Apresentar para visitantes, recolher as fichas do visitante e escrever, em grupo, o que aprenderam.', 'Apresentação e reflexão final')],
    ideias=[
        ('Separador de recicláveis', 'Treinar um modelo com fotos de papel, plástico e metal e testar com o lixo da sala. Pergunta: quantas fotos de cada tipo são precisas para o modelo acertar?'),
        ('Folhas da horta', 'Classificar folhas saudáveis e com manchas a partir de fotos. Pergunta: a luz da foto muda o resultado?'),
        ('Previsão do tempo da escola', 'Medir temperatura e nuvens por duas semanas numa planilha e criar uma regra de previsão. Pergunta: a regra acerta mais do que chutar?'),
        ('Merenda sem desperdício', 'Registrar as sobras por cardápio e descobrir quais pratos sobram menos. Pergunta: o que a turma pode sugerir à cozinha?'),
        ('Sons da sala', 'Treinar um modelo para reconhecer palmas, assobio e silêncio. Pergunta: o barulho de fundo atrapalha?'),
        ('Gestos que comandam', 'Treinar um modelo com gestos da mão para controlar uma apresentação. Pergunta: funciona igual com mãos de pessoas diferentes?'),
        ('Recomendador da biblioteca', 'Criar regras ou uma tabela que sugere livros a partir dos gostos dos colegas. Pergunta: as sugestões agradam?'),
        ('Gerador de imagens e estereótipos', 'Pedir imagens de profissões a um gerador e contar quem aparece. Pergunta: o que isso mostra sobre os dados de treino?'),
        ('Checador de boatos', 'Montar uma lista de sinais de notícia falsa e testar com notícias reais e falsas. Pergunta: quantos sinais bastam?'),
        ('Mapa do bairro', 'Coletar com formulário e foto pontos com buraco, lixo ou falta de luz e mostrar num mapa. Pergunta: onde a prefeitura deveria começar?'),
        ('Humano ou máquina', 'Comparar textos escritos por colegas e por um chatbot e ver se os visitantes acertam quem escreveu. Pergunta: o que entrega um texto de máquina?'),
        ('Energia na escola', 'Medir quanto tempo luzes e aparelhos ficam ligados e estimar o consumo. Pergunta: qual hábito economiza mais?')],
    banner=['Título curto e a pergunta do projeto', 'Por que o tema importa para a escola ou o bairro', 'Como a solução funciona (um desenho do caminho: dados, treino, teste)',
            'Os dados: quantos, de onde e como foram coletados', 'Resultados: o que acertou, o que melhorou entre as versões, um gráfico', 'Limites e cuidados: quando a IA erra e por quê',
            'Próximos passos e quem é o grupo'],
    etica=['O projeto usa fotos, vozes ou dados de pessoas? Todas autorizaram? Se forem menores de idade, os responsáveis também?', 'Os dados guardados identificam alguém (nome, rosto, endereço)? Dá para usar sem identificar?',
           'Onde os dados ficam guardados e quem tem acesso? Serão apagados depois da feira?', 'A ferramenta usada pede cadastro? Ele é permitido para a idade da turma, segundo os termos de uso?',
           'O resultado pode prejudicar ou constranger alguém se estiver errado?', 'O banner explica que a IA pode errar e mostra um exemplo de erro?'],
    rubrica=[('Pergunta e problema', 'Pergunta clara, ligada a um problema real e possível de testar.', 'Pergunta clara, mas ampla ou difícil de testar.', 'Tema escolhido, ainda sem pergunta.'),
             ('Dados', 'Dados suficientes, coletados com método e com autorização.', 'Poucos dados ou método pouco explicado.', 'Dados soltos, sem registro de como foram coletados.'),
             ('Solução e testes', 'Protótipo funcionando, testado com dados novos e melhorado.', 'Protótipo funcionando, com poucos testes.', 'Ideia descrita, protótipo incompleto.'),
             ('Ética e limites', 'Explica erros, vieses e cuidados com dados pessoais.', 'Cita limites de forma geral.', 'Não menciona limites nem cuidados.'),
             ('Comunicação', 'Banner organizado e fala clara; responde às perguntas dos visitantes.', 'Banner completo, fala lida.', 'Banner incompleto ou fala confusa.'),
             ('Trabalho em grupo', 'Diário de bordo em dia; todos participam e sabem explicar o projeto.', 'Diário com lacunas; participação desigual.', 'Diário vazio; uma pessoa faz quase tudo.')],
)

# ---------------------------------------------------------------- conteúdo: guia de adaptação
GUIA = dict(
    titulo='Guia de adaptação dos kits',
    lead='Como usar os jogos, as fichas e as avaliações da área premium com toda a turma, incluindo estudantes com deficiência, transtornos de aprendizagem e outras necessidades. As orientações seguem o Desenho Universal para a Aprendizagem (DUA): planejar desde o início caminhos diferentes para todos chegarem ao mesmo objetivo.',
    dua=[('Várias formas de apresentar', 'A mesma ideia em texto, imagem, fala, objeto concreto e jogo. Os kits já trazem o jogo, a versão desplugada, os slides e o cartaz: combine pelo menos dois.'),
         ('Várias formas de participar', 'Escolha de papéis no grupo, desafios em níveis e metas curtas. O placar de estrelas do painel da aula ajuda a mostrar o que cada dupla conquistou.'),
         ('Várias formas de mostrar o que aprendeu', 'Escrever, desenhar, falar, montar ou apontar. A versão acessível da avaliação aceita resposta oral e desenho na questão aberta.')],
    perfis=[
        ('Transtorno do espectro autista (TEA)', ['Mostre o roteiro visual da aula (versão acessível do kit) no início e risque cada etapa ao terminar.', 'Avise antes de mudanças, como troca de duplas ou de sala; o sorteio de duplas do painel pode ser substituído por duplas fixas.',
            'Dê instruções curtas, uma de cada vez, e confira o entendimento com uma pergunta direta.', 'Ofereça um lugar mais calmo e permita fones abafadores; desligue o som do jogo se incomodar.'],
         'Prefira a ficha com menos perguntas e aceite respostas curtas ou por escolha entre opções.'),
        ('TDAH', ['Divida a aula em blocos de 10 a 15 minutos; o cronômetro do painel da aula marca cada bloco.', 'Dê um papel ativo no grupo (quem clica, quem anota, quem cronometra) e troque os papéis.',
            'Mantenha à vista só o material da etapa atual.', 'Elogie o progresso de cada fase do jogo, não só o resultado final.'], 'Avaliação em partes, com pausas, e tempo extra.'),
        ('Dislexia', ['Use a versão acessível (fonte Lexend, letra ampliada e espaçamento maior).', 'Leia em voz alta os enunciados dos jogos e das fichas ou peça a um colega que leia.',
            'Evite copiar do quadro: entregue a ficha impressa ou o arquivo em Word.', 'Valorize a explicação oral do raciocínio.'], 'Leitura das questões em voz alta e correção pelo conteúdo, não pela ortografia.'),
        ('Deficiência intelectual', ['Comece pelo concreto: a versão desplugada com objetos e cartões antes do jogo no computador.', 'Reduza o número de fases e de perguntas, mantendo a mesma ideia central do mês.',
            'Use exemplos do dia a dia da turma e repita a ideia principal com outras palavras.', 'Trabalhe em dupla com um colega que explique com calma.'], 'Poucas questões, com apoio de imagens e escolha entre duas opções.'),
        ('Baixa visão', ['Imprima a versão acessível em A4 ou amplie para A3; mantenha o alto contraste (preto no branco).', 'No computador, use o zoom do navegador (Ctrl e +) e o modo escuro, se ajudar.',
            'Sente o estudante perto do projetor e envie os slides antes da aula.', 'Descreva em voz alta o que aparece na tela durante o jogo.'], 'Prova ampliada, tempo extra e boa iluminação.'),
        ('Cegueira', ['Use a versão acessível em Word: os títulos permitem navegar com o leitor de tela (NVDA no Windows, gratuito; TalkBack no Android).', 'Faça a versão desplugada com objetos em relevo, como cartões com texturas, tampinhas e barbante.',
            'Forme dupla com um colega que descreve a tela e deixe o estudante decidir as jogadas.', 'Descreva gráficos e imagens com palavras: o que mostram e qual é a ideia principal.'], 'Avaliação em Word com o leitor de tela ou oral.'),
        ('Surdez e deficiência auditiva', ['Garanta a presença do intérprete de Libras quando houver e entregue antes os termos novos da aula.', 'Use a tradução para Libras do VLibras (gratuito) nos textos digitais.',
            'Mantenha as instruções escritas no quadro ou no painel da aula e evite falar de costas para a turma.', 'Prefira jogos com retorno visual; desative a dependência de som.'], 'Enunciados claros e curtos; aceite respostas em Libras gravadas em vídeo.'),
        ('Deficiência física ou motora', ['Verifique se o estudante consegue usar o mouse; vários jogos funcionam também pelo teclado e pela tela de toque.', 'Dê mais tempo nas fases com tempo e arraste.',
            'Na versão desplugada, use peças maiores e uma mesa na altura adequada.', 'Deixe um colega executar os cliques seguindo as decisões do estudante.'], 'Resposta oral ou marcada por um colega, conforme a indicação do estudante.')],
    conferir=['O objetivo da aula está escrito em uma frase simples e visível?', 'A ideia aparece em pelo menos duas formas (jogo e desplugado, slide e cartaz)?', 'Há um roteiro visual com as etapas da aula?',
              'Os materiais impressos têm letra de pelo menos 14 pontos e bom contraste?', 'As imagens importantes têm descrição em palavras?', 'Existe mais de um jeito de responder (escrever, falar, desenhar)?',
              'As duplas e os papéis foram pensados com antecedência?', 'O tempo previsto inclui pausas e tempo extra?', 'Conversei com o AEE ou com a família sobre o que funciona para cada estudante?'],
    recursos=[('NVDA', 'leitor de tela gratuito para Windows (nvaccess.org)'), ('VLibras', 'tradução de textos digitais para Libras, gratuito (gov.br/governodigital)'),
              ('Lupa e alto contraste', 'já vêm no Windows, no Android e no iPhone, nas configurações de acessibilidade'), ('Fonte Lexend', 'fonte gratuita pensada para facilitar a leitura (lexend.com)')],
)


# ---------------------------------------------------------------- HTML para PDF
def pagina(titulo, corpo):
    return (f'<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><title>{e(titulo)}</title><meta name="author" content="{AUTORA}"><style>{CSS}'
            '.ideias { display: grid; grid-template-columns: 1fr 1fr; gap: 8px; } .ideias div { border: 1.5px solid #111; border-radius: 14px; padding: 8px 12px; font-size: 10pt; break-inside: avoid; }'
            '.ideias b { display: block; font-size: 11pt; } .cert { height: 250mm; border: 4px solid #111; border-radius: 24px; padding: 30mm 20mm; text-align: center; display: flex; flex-direction: column; justify-content: center; }'
            '.cert h1 { font-size: 40pt; } .cert p { font-size: 15pt; } .perfil { break-inside: avoid; border: 1.5px solid #111; border-radius: 16px; padding: 8px 14px; margin: 10px 0; }'
            '.perfil h2 { margin: 4px 0; } .caixa-v { border-left: 6px solid #79c257; padding: 4px 12px; margin: 6px 0; } .box { width: 16px; height: 16px; border: 1.5px solid #111; display: inline-block; vertical-align: middle; margin-right: 8px; }'
            f'</style></head><body><div class="faixa"><span>ohtopico · Área premium</span><b>{e(titulo)}</b></div>{corpo}</body></html>')


def feira_html(f):
    lis = lambda xs: ''.join(f'<li>{e(x)}</li>' for x in xs)
    corpo = (f'<h1>{e(f["titulo"])}</h1><p class="sub">{e(f["lead"])}</p>'
             f'<div class="caixas"><span class="caixa">6 semanas, 1 ou 2 aulas por semana</span><span class="caixa">Grupos de 3 a 5</span><span class="caixa">BNCC: {", ".join(c for c, _ in f["bncc"])}</span></div>'
             '<h2>Cronograma</h2><table><tr><th>Semana</th><th>Etapa</th><th>O que fazer</th><th>Entrega</th></tr>' +
             ''.join(f'<tr><td>{s}</td><td><b>{e(t)}</b></td><td>{e(x)}</td><td>{e(en)}</td></tr>' for s, t, x, en in f['semanas']) + '</table>'
             '<h2>12 ideias de projeto</h2><p>Os grupos podem escolher uma ideia, adaptar ou propor outra. Ferramentas que funcionam no navegador sem cadastro de estudante, ou a versão desplugada com cartões, servem para todas.</p>'
             '<div class="ideias">' + ''.join(f'<div><b>{i}. {e(t)}</b>{e(x)}</div>' for i, (t, x) in enumerate(f['ideias'], 1)) + '</div>'
             '<div class="quebra"></div><h1 style="font-size:22pt">Ficha de inscrição do grupo</h1>'
             '<div class="nome"><span>Nome do projeto:</span><span>Turma:</span><span>Data:</span></div>'
             '<div class="q"><p>Integrantes e papéis (coordenação, dados, protótipo, banner, apresentação)</p>' + '<div class="linha"></div>' * 5 + '</div>'
             '<div class="q"><p>Nossa pergunta</p>' + '<div class="linha"></div>' * 2 + '</div>'
             '<div class="q"><p>Por que esse tema importa para a escola ou o bairro</p>' + '<div class="linha"></div>' * 3 + '</div>'
             '<div class="q"><p>Que dados vamos usar e como vamos coletar</p>' + '<div class="linha"></div>' * 3 + '</div>'
             '<div class="q"><p>Material ou ferramenta de que vamos precisar</p>' + '<div class="linha"></div>' * 2 + '</div>'
             '<div class="quebra"></div><h1 style="font-size:22pt">Diário de bordo</h1><p class="sub">Uma página por semana. Comece sempre pelo que deu certo.</p>' +
             ''.join(f'<div class="q"><p>Semana {s} · {e(t)}</p><table><tr><th style="width:34%">O que deu certo</th><th style="width:33%">O que foi difícil</th><th>Próximo passo</th></tr><tr><td style="height:70px"></td><td></td><td></td></tr></table></div>'
                     for s, t, _, _ in f['semanas']) +
             '<div class="quebra"></div><h1 style="font-size:22pt">Roteiro do banner</h1><ol class="passos">' + lis(f['banner']) + '</ol>'
             '<h2>Cuidados com dados e ética</h2><p>Antes de coletar e antes da feira, o grupo responde:</p><ul>' + ''.join(f'<li><span class="box"></span>{e(x)}</li>' for x in f['etica']) + '</ul>'
             '<div class="quebra"></div><h1 style="font-size:22pt">Rubrica de avaliação</h1><table><tr><th style="width:18%">Critério</th><th>Completo (3)</th><th>Em construção (2)</th><th>Inicial (1)</th></tr>' +
             ''.join(f'<tr><td><b>{e(c)}</b></td><td>{e(a)}</td><td>{e(b)}</td><td>{e(d)}</td></tr>' for c, a, b, d in f['rubrica']) + '</table>'
             '<p class="sub">Some os pontos (6 a 18). Use a mesma rubrica para a autoavaliação do grupo na semana 5, comparando com a avaliação final.</p>'
             '<h2>Ficha do visitante</h2><p>Imprima várias por página e deixe junto ao banner.</p>' +
             ''.join('<div class="leve"><b>Projeto:</b> ______________________________<br>O que eu mais gostei: ______________________________________<br>'
                     'Uma pergunta para o grupo: ___________________________________<br>Eu aprendi que: ______________________________________________</div>' for _ in range(3)) +
             '<div class="quebra"></div><div class="cert"><p>ohtopico · Feira de Ciências com IA</p><h1>Certificado</h1><p>Certificamos que</p><p>__________________________________________</p>'
             '<p>participou da Feira de Ciências com IA com o projeto</p><p>__________________________________________</p><p>Escola: ____________________________</p><p>Data: ___/___/_____</p>'
             '<p style="margin-top:20mm">______________________________<br>Professor(a) responsável</p></div>')
    return pagina(f['titulo'], corpo)


def guia_html(g):
    corpo = (f'<h1>{e(g["titulo"])}</h1><p class="sub">{e(g["lead"])}</p><h2>Os três caminhos do DUA</h2>' +
             ''.join(f'<div class="caixa-v"><b>{e(t)}.</b> {e(x)}</div>' for t, x in g['dua']) +
             '<h2>O que cada kit já oferece</h2><ul><li><b>Versão acessível</b> (PDF e Word): roteiro visual da aula, ficha e avaliação em letra ampliada e alto contraste, com resposta oral ou por desenho.</li>'
             '<li><b>Versão desplugada</b>: a mesma ideia com objetos e cartões, sem tela.</li><li><b>Painel da aula</b>: cronômetro, etapas e duplas na tela, para antecipar e marcar o tempo.</li>'
             '<li><b>Arquivos editáveis</b>: para reduzir, trocar exemplos e ampliar o que for preciso.</li></ul>' +
             ''.join(f'<div class="perfil"><h2>{e(t)}</h2><ul>{"".join(f"<li>{e(x)}</li>" for x in xs)}</ul><p><b>Na avaliação:</b> {e(av)}</p></div>' for t, xs, av in g['perfis']) +
             '<h2>Lista de conferência para planejar a aula</h2><ul>' + ''.join(f'<li><span class="box"></span>{e(x)}</li>' for x in g['conferir']) + '</ul>'
             '<h2>Recursos gratuitos</h2><ul>' + ''.join(f'<li><b>{e(n)}</b>: {e(x)}</li>' for n, x in g['recursos']) + '</ul>'
             '<p class="sub">Cada estudante é único: estas orientações são um ponto de partida. Converse com o atendimento educacional especializado (AEE), com a família e com o próprio estudante sobre o que funciona melhor.</p>')
    return pagina(g['titulo'], corpo)


# ---------------------------------------------------------------- Word
def d_vazio(titulo):
    return {'nome_mes': '', 'titulo': titulo}


def feira_docx(f, destino):
    doc = word(d_vazio(f['titulo'])); doc.add_heading(f['titulo'], level=0); doc.add_paragraph(f['lead'])
    doc.add_paragraph('BNCC: ' + '; '.join(f'{c} {x}' for c, x in f['bncc']))
    doc.add_heading('Cronograma', 1); tabela(doc, [['Semana', 'Etapa', 'O que fazer', 'Entrega']] + [list(x) for x in f['semanas']])
    doc.add_heading('12 ideias de projeto', 1)
    for t, x in f['ideias']:
        p = doc.add_paragraph(style='List Number'); p.add_run(t + '. ').bold = True; p.add_run(x)
    doc.add_page_break(); doc.add_heading('Ficha de inscrição do grupo', 1)
    doc.add_paragraph('Nome do projeto: ____________________________   Turma: ________   Data: ___/___/_____')
    for q, n in (('Integrantes e papéis (coordenação, dados, protótipo, banner, apresentação)', 5), ('Nossa pergunta', 2), ('Por que esse tema importa para a escola ou o bairro', 3),
                 ('Que dados vamos usar e como vamos coletar', 3), ('Material ou ferramenta de que vamos precisar', 2)):
        doc.add_paragraph().add_run(q).bold = True; linhas(doc, n)
    doc.add_page_break(); doc.add_heading('Diário de bordo', 1); doc.add_paragraph('Uma página por semana. Comece sempre pelo que deu certo.')
    for s, t, _, _ in f['semanas']:
        doc.add_heading(f'Semana {s} · {t}', 2); tabela(doc, [['O que deu certo', 'O que foi difícil', 'Próximo passo'], ['\n\n\n', '', '']])
    doc.add_page_break(); doc.add_heading('Roteiro do banner', 1)
    for x in f['banner']:
        doc.add_paragraph(x, style='List Number')
    doc.add_heading('Cuidados com dados e ética', 1)
    for x in f['etica']:
        doc.add_paragraph('[  ] ' + x)
    doc.add_page_break(); doc.add_heading('Rubrica de avaliação', 1)
    tabela(doc, [['Critério', 'Completo (3)', 'Em construção (2)', 'Inicial (1)']] + [list(x) for x in f['rubrica']])
    doc.add_paragraph('Some os pontos (6 a 18). Use a mesma rubrica para a autoavaliação do grupo na semana 5.')
    doc.add_heading('Ficha do visitante', 1)
    for _ in range(2):
        doc.add_paragraph('Projeto: ______________________________\nO que eu mais gostei: ______________________________\nUma pergunta para o grupo: ______________________________\nEu aprendi que: ______________________________')
    doc.add_page_break(); h = doc.add_heading('Certificado', 0)
    for x in ('Certificamos que ______________________________________', 'participou da Feira de Ciências com IA com o projeto ______________________________',
              'Escola: ____________________________   Data: ___/___/_____', '\n\n______________________________\nProfessor(a) responsável'):
        r = doc.add_paragraph().add_run(x); r.font.size = Pt(16)
    rodape(doc); doc.save(destino)


def guia_docx(g, destino):
    doc = word(d_vazio(g['titulo'])); doc.add_heading(g['titulo'], level=0); doc.add_paragraph(g['lead'])
    doc.add_heading('Os três caminhos do DUA', 1)
    for t, x in g['dua']:
        p = doc.add_paragraph(style='List Bullet'); p.add_run(t + '. ').bold = True; p.add_run(x)
    for t, xs, av in g['perfis']:
        doc.add_heading(t, 1)
        for x in xs:
            doc.add_paragraph(x, style='List Bullet')
        p = doc.add_paragraph(); p.add_run('Na avaliação: ').bold = True; p.add_run(av)
    doc.add_heading('Lista de conferência para planejar a aula', 1)
    for x in g['conferir']:
        doc.add_paragraph('[  ] ' + x)
    doc.add_heading('Recursos gratuitos', 1)
    for n, x in g['recursos']:
        p = doc.add_paragraph(style='List Bullet'); p.add_run(n + ': ').bold = True; p.add_run(x)
    rodape(doc); doc.save(destino)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--fontes', default='_privado/premium'); ap.add_argument('--tipos', required=True); a = ap.parse_args()
    arq = Path(a.fontes) / 'arquivos'
    (arq / 'imprimir').mkdir(parents=True, exist_ok=True); (arq / 'editar').mkdir(parents=True, exist_ok=True)
    imprimir([(feira_html(FEIRA), arq / 'imprimir' / 'feira-de-ciencias-com-ia.pdf'), (guia_html(GUIA), arq / 'imprimir' / 'guia-de-adaptacao-dos-kits.pdf')], a.tipos)
    feira_docx(FEIRA, arq / 'editar' / 'feira-de-ciencias-com-ia.docx'); guia_docx(GUIA, arq / 'editar' / 'guia-de-adaptacao-dos-kits.docx')
    nv = arq / 'niveis.json'; n = json.loads(nv.read_text()); n.update(imprimir=2, editar=3); nv.write_text(json.dumps(n))
    print('ok extras')


if __name__ == '__main__':
    main()
